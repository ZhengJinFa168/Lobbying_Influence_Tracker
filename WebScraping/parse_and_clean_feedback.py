#!/usr/bin/env python3
"""
parse_and_clean_feedback.py
===========================
Script unico e parametrico per il parsing, l'estrazione multi-formato,
la pulizia del testo legale, il rilevamento di riferimenti normativi (regex)
e il chunking semantico (per modelli di embedding) dei feedback Have Your Say.

Supporta qualsiasi procedura legislativa definita in data_sources/procedures/<id>.yaml (default: ai_act):
  - Stage 1: Inception Impact Assessment
  - Stage 2: Proposal for a Regulation

Autore: Antigravity AI (Senior Data Engineer)
"""

import os
import sys
import json
import re
import unicodedata
import argparse
from typing import List, Dict, Any, Optional, Tuple
import yaml

# Librerie di estrazione PDF e DOCX
try:
    from pypdf import PdfReader
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

try:
    import pdfplumber
    PDFPLUMBER_AVAILABLE = True
except ImportError:
    PDFPLUMBER_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from tqdm import tqdm
    TQDM_AVAILABLE = True
except ImportError:
    TQDM_AVAILABLE = False


# =====================================================================
# CONFIGURAZIONE STAGES
# =====================================================================

def load_stages_config(procedure_id: str = "ai_act", base_dir: Optional[str] = None) -> Dict[int, Dict[str, Any]]:
    """
    Carica la configurazione degli stage da data_sources/procedures/<procedure_id>.yaml.
    Mappa:
      - consultation.A.publication -> Stage 1 ("inception", cartella feedback/stage1_inception_<pub>/)
      - consultation.B.publication -> Stage 2 ("proposal", cartella feedback/stage2_proposal_<pub>/)
    I track la cui pubblicazione è null vengono ignorati.
    """
    if base_dir is None:
        base_dir = os.path.dirname(os.path.abspath(__file__))
    repo_root = os.path.abspath(os.path.join(base_dir, ".."))
    proc_id = procedure_id.strip().removesuffix(".yaml")
    procedures_dir = os.path.join(repo_root, "data_sources", "procedures")
    yaml_path = os.path.join(procedures_dir, f"{proc_id}.yaml")

    if not os.path.isfile(yaml_path):
        known = ", ".join(sorted(f.removesuffix(".yaml") for f in os.listdir(procedures_dir) if f.endswith(".yaml")))
        raise SystemExit(f"unknown procedure {proc_id!r}: expected data_sources/procedures/<id>.yaml, one of: {known}")

    with open(yaml_path, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}

    consultation = data.get("consultation", {})
    configs: Dict[int, Dict[str, Any]] = {}

    # Stage 1: Track A (inception)
    cA = consultation.get("A")
    if isinstance(cA, dict) and cA.get("publication") is not None:
        pub_a = int(cA["publication"])
        folder_a = f"feedback/stage1_inception_{pub_a}"
        if proc_id == "ai_act":
            title_a = "Stage 1 (Inception Impact Assessment 2020 - PubID 13340)"
        else:
            label_a = (cA.get("label") or "Inception Impact Assessment").strip()
            title_a = f"Stage 1 ({label_a} - PubID {pub_a})"

        configs[1] = {
            "publication_id": pub_a,
            "stage_name": "STAGE_1_INCEPTION_FEEDBACK",
            "title": title_a,
            "folder": folder_a,
            "output_file": f"{folder_a}/raw_feedback_stage1_{pub_a}.json",
            "raw_json": f"{folder_a}/raw_feedback_stage1_{pub_a}.json",
            "parsed_json": f"{folder_a}/parsed_feedback_stage1_{pub_a}.json",
            "attachments_dir": f"{folder_a}/attachments",
        }

    # Stage 2: Track B (proposal)
    cB = consultation.get("B")
    if isinstance(cB, dict) and cB.get("publication") is not None:
        pub_b = int(cB["publication"])
        folder_b = f"feedback/stage2_proposal_{pub_b}"
        if proc_id == "ai_act":
            title_b = "Stage 2 (Proposal for a Regulation 2021 - PubID 14488)"
        else:
            label_b = (cB.get("label") or "Proposal for a Regulation").strip()
            title_b = f"Stage 2 ({label_b} - PubID {pub_b})"

        configs[2] = {
            "publication_id": pub_b,
            "stage_name": "STAGE_2_PROPOSAL_FEEDBACK",
            "title": title_b,
            "folder": folder_b,
            "output_file": f"{folder_b}/raw_feedback_stage2_{pub_b}.json",
            "raw_json": f"{folder_b}/raw_feedback_stage2_{pub_b}.json",
            "parsed_json": f"{folder_b}/parsed_feedback_stage2_{pub_b}.json",
            "attachments_dir": f"{folder_b}/attachments",
        }

    return configs


# Configurazione di default per retrocompatibilità (ai_act)
STAGES_CONFIG = load_stages_config("ai_act")


# =====================================================================
# 1. ESTRAZIONE MULTI-FORMATO ROBUSTA CON MAGIC BYTES
# =====================================================================

def extract_text_from_file_robust(file_path: str) -> Tuple[List[str], bool, Optional[str]]:
    """
    Estrae il testo da PDF o DOCX determinando il formato reale dai 'magic bytes'
    del file (risolve i casi in cui l'API della Commissione converte i file DOCX in PDF).
    
    Ritorna:
        (pages_or_paragraphs: List[str], has_text: bool, error_msg: Optional[str])
    """
    if not os.path.isfile(file_path):
        return [], False, f"File non trovato: {file_path}"

    try:
        with open(file_path, "rb") as f:
            magic = f.read(5)
    except Exception as e:
        return [], False, f"Impossibile leggere file: {e}"

    pages_text: List[str] = []
    error_msg: Optional[str] = None

    # Caso 1: È un file PDF (inizia con %PDF)
    if magic.startswith(b"%PDF"):
        # Tentativo 1: pypdf
        if PYPDF_AVAILABLE:
            try:
                reader = PdfReader(file_path)
                for page in reader.pages:
                    try:
                        text = page.extract_text() or ""
                        pages_text.append(text)
                    except Exception:
                        pages_text.append("")
            except Exception as e:
                error_msg = f"pypdf error: {e}"
                pages_text = []

        total_len = sum(len(p.strip()) for p in pages_text)

        # Tentativo 2: pdfplumber se pypdf ha estratto meno di 50 caratteri
        if total_len < 50 and PDFPLUMBER_AVAILABLE:
            try:
                plumber_pages = []
                with pdfplumber.open(file_path) as pdf:
                    for page in pdf.pages:
                        try:
                            text = page.extract_text() or ""
                            plumber_pages.append(text)
                        except Exception:
                            plumber_pages.append("")
                plumber_len = sum(len(p.strip()) for p in plumber_pages)
                if plumber_len > total_len:
                    pages_text = plumber_pages
                    total_len = plumber_len
                    error_msg = None
            except Exception as e:
                if not error_msg:
                    error_msg = f"pdfplumber error: {e}"

        return pages_text, total_len >= 30, error_msg

    # Caso 2: È un archivio DOCX (inizia con PK\x03\x04)
    elif magic.startswith(b"PK") and DOCX_AVAILABLE:
        try:
            doc = docx.Document(file_path)
            paragraphs = []
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt:
                    paragraphs.append(txt)
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        paragraphs.append(" | ".join(cells))
            total_len = sum(len(p) for p in paragraphs)
            return paragraphs, total_len >= 30, None
        except Exception as e:
            return [], False, f"docx error: {e}"

    return [], False, "Formato file non riconosciuto o libreria non disponibile"


def resolve_attachment_path(att_dict: Dict[str, Any], attachments_dir: str) -> Tuple[Optional[str], Optional[str]]:
    """
    Risolve il percorso assoluto e il nome file effettivo di un allegato sul filesystem.
    """
    file_name = att_dict.get("fileName")
    ers_file_name = att_dict.get("ersFileName")
    local_path = att_dict.get("localFilePath")

    candidates = []
    if local_path:
        candidates.append(local_path)
        candidates.append(os.path.join(attachments_dir, os.path.basename(local_path)))
    if file_name:
        candidates.append(os.path.join(attachments_dir, file_name))
    if ers_file_name:
        candidates.append(os.path.join(attachments_dir, ers_file_name))

    for cand in candidates:
        if cand and os.path.isfile(cand):
            return os.path.abspath(cand), os.path.basename(cand)

    # Ricerca case-insensitive
    if os.path.isdir(attachments_dir) and (file_name or ers_file_name):
        t_name = (file_name or ers_file_name).lower()
        t_base = os.path.splitext(t_name)[0]
        for existing in os.listdir(attachments_dir):
            e_lower = existing.lower()
            if e_lower == t_name or os.path.splitext(e_lower)[0] == t_base:
                full_p = os.path.join(attachments_dir, existing)
                return os.path.abspath(full_p), existing

    chosen_name = file_name or ers_file_name or (os.path.basename(local_path) if local_path else None)
    return None, chosen_name


# =====================================================================
# 2. PULIZIA E SANIFICAZIONE AVANZATA DEL TESTO LEGALE
# =====================================================================

def clean_page_text(page_text: str) -> str:
    """
    Pulisce il testo di una singola pagina:
    - Normalizzazione Unicode NFKC
    - Sostituzione di ligature tipografiche
    - Rimozione timbri Ares e numeri di pagina
    """
    if not page_text:
        return ""

    text = unicodedata.normalize("NFKC", page_text)

    replacements = {
        "ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl",
        "\u00a0": " ", "\u200b": "", "\ufeff": "", "\r": "",
        "“": '"', "”": '"', "„": '"', "«": '"', "»": '"',
        "‘": "'", "’": "'", "ʻ": "'", "ʼ": "'",
        "–": "-", "—": "-", "―": "-", "\t": " "
    }
    for k, v in replacements.items():
        text = text.replace(k, v)

    lines = [line.strip() for line in text.split("\n")]
    cleaned_lines = []

    for line in lines:
        if not line:
            cleaned_lines.append("")
            continue

        # Numeri di pagina
        if re.match(r"^(?:page|p\.)\s*\d+\s*(?:of|/)\s*\d+$", line, re.IGNORECASE):
            continue
        if re.match(r"^(?:page|p\.)\s*\d+$", line, re.IGNORECASE):
            continue
        if re.match(r"^[-\u2013\u2014]\s*\d+\s*[-\u2013\u2014]$", line):
            continue
        if re.match(r"^\d{1,3}$", line):
            continue

        # Timbro Ares
        if re.match(r"^Ref\.\s*Ares\(\d+\)\d+\s*-\s*\d{2}/\d{2}/\d{4}$", line):
            continue

        cleaned_lines.append(line)

    return "\n".join(cleaned_lines)


def filter_repetitive_headers_footers(pages_text: List[str]) -> List[str]:
    """
    Rileva e rimuove header o footer che si ripetono identici su oltre il 40% delle pagine.
    """
    if len(pages_text) <= 2:
        return pages_text

    header_counts: Dict[str, int] = {}
    footer_counts: Dict[str, int] = {}

    for page in pages_text:
        lines = [l.strip() for l in page.split("\n") if l.strip()]
        if lines:
            if len(lines[0]) < 120:
                header_counts[lines[0]] = header_counts.get(lines[0], 0) + 1
            if len(lines[-1]) < 120:
                footer_counts[lines[-1]] = footer_counts.get(lines[-1], 0) + 1

    threshold = max(2, int(len(pages_text) * 0.4))
    rep_headers = {k for k, v in header_counts.items() if v >= threshold}
    rep_footers = {k for k, v in footer_counts.items() if v >= threshold}

    result = []
    for page in pages_text:
        lines = [l.strip() for l in page.split("\n")]
        valid = [l for l in lines if l not in rep_headers and l not in rep_footers]
        result.append("\n".join(valid))

    return result


def sanitize_extracted_text(text: str) -> str:
    """
    Sanifica e normalizza il testo estratto:
    1. Rimuove sequenze di pallini, trattini, asterischi ripetuti (es. • • • • o ------ o .....)
    2. Ripara parole spezzate con trattino (es. 'future -proof' -> 'future-proof', 'high -risk' -> 'high-risk')
    3. Rimuove caratteri di controllo non stampabili
    4. Ripara spazi errati prima della punteggiatura (es. 'technology , we' -> 'technology, we')
    5. Normalizza spazi multipli e troppi a capo consecutivi
    6. Rimuove righe isolate prive di contenuto alfanumerico
    """
    if not text:
        return ""

    # 1. Rimuove sequenze di pallini, trattini, asterischi o puntini decorativi ripetuti (3 o più)
    text = re.sub(r'([•·\-\*\._~=]\s*){3,}', ' ', text)

    # 2. Ripara parole spezzate con trattino (senza intaccare i trattini lunghi tra proposizioni)
    text = re.sub(r'(\b\w+)\s+-(?=[^\s\-])', r'\1-', text)
    text = re.sub(r'(?<=[^\s\-])-\s+(\w+\b)', r'-\1', text)

    # 3. Rimuove caratteri di controllo non stampabili (tranne newline e tab)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', text)

    # 4. Ripara spazi spuri prima di punteggiatura
    text = re.sub(r'\s+([,;:?.!])', r'\1', text)

    # 5. Normalizza spazi multipli orizzontali
    text = re.sub(r'[ \t]+', ' ', text)

    # 6. Normalizza troppi a capo consecutivi
    text = re.sub(r'\n{3,}', '\n\n', text)

    # 7. Rimuove righe prive di contenuto alfanumerico
    lines = [line.strip() for line in text.split('\n')]
    cleaned_lines = []
    for line in lines:
        if not line:
            cleaned_lines.append('')
            continue
        if not re.search(r'[a-zA-Z0-9]', line):
            continue
        cleaned_lines.append(line)

    joined = '\n'.join(cleaned_lines)
    joined = re.sub(r'\n{3,}', '\n\n', joined)
    return joined.strip()


def postprocess_legal_text(text: str) -> str:
    """
    Ricompatta i paragrafi, ricompone le parole sillabate a fine riga
    e applica la sanificazione finale.
    """
    if not text or not text.strip():
        return ""

    # Ricomposizione parole sillabate (es. "regu-\nlation" -> "regulation")
    text = re.sub(r'(\b[a-zA-Z]{2,})-\s*\n\s*([a-zA-Z]{2,}\b)', r'\1\2', text)

    raw_paragraphs = re.split(r'\n\s*\n+', text)
    cleaned_paragraphs = []

    for p in raw_paragraphs:
        p = p.strip()
        if not p:
            continue

        lines = p.split('\n')
        processed_lines = []
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if re.match(r'^([•\-\*■\u2022]|\d+[\.\)]|[a-zA-Z][\.\)])\s+', line):
                processed_lines.append('\n' + line)
            else:
                processed_lines.append(line)

        joined = ' '.join(processed_lines)
        joined = re.sub(r'\s*\n\s*', '\n', joined)
        joined = re.sub(r'[ \t]+', ' ', joined)
        cleaned_paragraphs.append(joined.strip())

    combined = '\n\n'.join(cleaned_paragraphs).strip()
    return sanitize_extracted_text(combined)


# =====================================================================
# 3. UNIONE CONTENUTO (FULL TEXT)
# =====================================================================

def is_placeholder_feedback(feedback: str) -> bool:
    """True se il commento web è solo un segnaposto per l'allegato."""
    if not feedback:
        return True
    f_lower = feedback.strip().lower()
    triggers = [
        "attached", "attachment", "enclosed", "find attached", "find enclosed",
        "see file", "see attached", "in the attached", "pleas find", "please see",
        "detailed input in the", "comments in the attached", "see the attached",
        "refer to our attached", "please refer to", "submission is attached"
    ]
    if len(f_lower) < 160 and any(kw in f_lower for kw in triggers):
        return True
    return False


def combine_feedback_and_attachment(web_feedback: str, attachment_text: str) -> str:
    """Combina in modo organico commento web ed eventuale allegato."""
    web_clean = postprocess_legal_text(clean_page_text(web_feedback)) if web_feedback else ""
    att_clean = attachment_text.strip() if attachment_text else ""

    if att_clean:
        if is_placeholder_feedback(web_clean):
            return att_clean
        if web_clean and len(web_clean) > 50 and (web_clean[:100].lower() in att_clean.lower()):
            return att_clean
        if web_clean:
            return f"{web_clean}\n\n{att_clean}"
        return att_clean
    return web_clean


# =====================================================================
# 4. ESTRAZIONE RIFERIMENTI NORMATIVI (REGEX)
# =====================================================================

def extract_legal_references(text: str) -> Dict[str, List[str]]:
    """Rileva e normalizza Articoli, Allegati e Considerando citati."""
    if not text:
        return {"detected_articles": [], "detected_annexes": [], "detected_recitals": []}

    articles = set()
    annexes = set()
    recitals = set()

    # Articoli singoli
    for match in re.finditer(r'\b(?:Article|Art\.|Articolo|Artikel)\s*([0-9]+[a-z]?(?:\s*\(\s*[0-9a-zA-Z]+\s*\))*)', text, re.IGNORECASE):
        art_ref = re.sub(r'\s+', '', match.group(1).strip())
        articles.add(f"Article {art_ref}")

    # Articoli multipli
    for match in re.finditer(r'\b(?:Articles|Articoli)\s+([0-9]+(?:\s*(?:,|and|or|to|\u2013|-)\s*[0-9]+)+)', text, re.IGNORECASE):
        nums = re.findall(r'\b[0-9]+\b', match.group(1))
        for n in nums:
            articles.add(f"Article {n}")

    # Allegati
    for match in re.finditer(r'\b(?:Annex|Annexes|Allegato|Anhang)\s+([IVXLCDM]+|[0-9]+|[A-E])\b', text, re.IGNORECASE):
        ann_ref = match.group(1).strip().upper()
        annexes.add(f"Annex {ann_ref}")

    for match in re.finditer(r'\b(?:Annexes|Allegati)\s+([IVXLCDM0-9A-E,\s\u2013\-andor]+)', text, re.IGNORECASE):
        tokens = re.findall(r'\b([IVXLCDM]+|[0-9]+|[A-E])\b', match.group(1), re.IGNORECASE)
        for t in tokens:
            annexes.add(f"Annex {t.upper()}")

    # Considerando
    for match in re.finditer(r'\b(?:Recital|Considerando|Erwägungsgrund)\s+([0-9]+)\b', text, re.IGNORECASE):
        rec_num = match.group(1).strip()
        recitals.add(f"Recital {rec_num}")

    for match in re.finditer(r'\b(?:Recitals|Considerandi)\s+([0-9]+(?:\s*(?:,|and|or|to|\u2013|-)\s*[0-9]+)+)', text, re.IGNORECASE):
        nums = re.findall(r'\b[0-9]+\b', match.group(1))
        for n in nums:
            recitals.add(f"Recital {n}")

    def natural_sort_key(s: str):
        return [int(t) if t.isdigit() else t.lower() for t in re.split(r'(\d+)', s)]

    return {
        "detected_articles": sorted(list(articles), key=natural_sort_key),
        "detected_annexes": sorted(list(annexes), key=natural_sort_key),
        "detected_recitals": sorted(list(recitals), key=natural_sort_key)
    }


# =====================================================================
# 5. CHUNKING INTELLIGENTE PER EMBEDDINGS
# =====================================================================

def chunk_legal_text(text: str, target_words: int = 400, overlap_words: int = 50) -> List[str]:
    """Divide il testo in chunk di 300-500 parole con overlap a livello di frase."""
    if not text or not text.strip():
        return []

    words = text.split()
    if len(words) <= target_words:
        return [text.strip()]

    sentences = re.split(r'(?<=[.!?])\s+', text)
    chunks = []
    current_chunk: List[str] = []
    current_word_count = 0

    for sentence in sentences:
        s_words = sentence.split()
        if not s_words:
            continue

        if len(s_words) > target_words:
            if current_chunk:
                chunks.append(" ".join(current_chunk).strip())
                current_chunk = []
                current_word_count = 0
            step = max(1, target_words - overlap_words)
            for i in range(0, len(s_words), step):
                chunk_slice = s_words[i : i + target_words]
                if chunk_slice:
                    chunks.append(" ".join(chunk_slice).strip())
            continue

        if current_word_count + len(s_words) > target_words and current_chunk:
            chunks.append(" ".join(current_chunk).strip())
            overlap_chunk = []
            overlap_count = 0
            for prev_s in reversed(current_chunk):
                p_len = len(prev_s.split())
                if overlap_count + p_len <= overlap_words or not overlap_chunk:
                    overlap_chunk.insert(0, prev_s)
                    overlap_count += p_len
                else:
                    break
            current_chunk = overlap_chunk
            current_word_count = overlap_count

        current_chunk.append(sentence)
        current_word_count += len(s_words)

    if current_chunk:
        final_chunk = " ".join(current_chunk).strip()
        if not chunks or final_chunk != chunks[-1]:
            chunks.append(final_chunk)

    return chunks


# =====================================================================
# 6. PIPELINE DI PARSING PER SINGOLO STAGE
# =====================================================================

def parse_stage(
    stage_num: int,
    base_dir: str,
    target_chunk_words: int = 400,
    overlap_words: int = 50,
    stages_config: Optional[Dict[int, Dict[str, Any]]] = None
):
    if stages_config is None:
        stages_config = STAGES_CONFIG
    conf = stages_config[stage_num]
    raw_path = os.path.join(base_dir, conf["raw_json"])
    parsed_path = os.path.join(base_dir, conf["parsed_json"])
    attachments_dir = os.path.join(base_dir, conf["attachments_dir"])

    if not os.path.isfile(raw_path):
        print(f"❌ File grezzo non trovato per {conf['title']}: {raw_path}")
        return

    with open(raw_path, "r", encoding="utf-8") as f:
        records = json.load(f)

    print("\n" + "=" * 70)
    print(f"🔍 AVVIO PARSING: {conf['title']}")
    print(f"📖 Record da elaborare: {len(records)}")
    print(f"📁 Cartella allegati: '{conf['attachments_dir']}'")
    print(f"💾 File di output: '{conf['parsed_json']}'")
    print("=" * 70 + "\n")

    parsed_records = []
    stats = {
        "total": len(records),
        "attachments_extracted": 0,
        "with_text": 0,
        "chunks_generated": 0,
        "articles_found": 0,
        "annexes_found": 0
    }

    iterator = enumerate(records, start=1)
    if TQDM_AVAILABLE:
        iterator = tqdm(iterator, total=len(records), desc=f"Stage {stage_num}")

    for idx, item in iterator:
        feedback_id = item.get("id")
        org_name = item.get("organization") or "Unknown"
        user_type = item.get("userType") or "UNKNOWN"
        country = item.get("country") or ""
        raw_web_feedback = item.get("feedback") or ""
        attachments_list = item.get("attachments") or []

        has_attachment_decl = len(attachments_list) > 0
        has_attachment_file = False
        attachment_name = None
        attachment_clean_text = ""
        extraction_error = None

        if has_attachment_decl:
            primary_att = attachments_list[0]
            resolved_p, attachment_name = resolve_attachment_path(primary_att, attachments_dir)

            if resolved_p:
                has_attachment_file = True
                stats["attachments_extracted"] += 1
                sections, has_txt, err = extract_text_from_file_robust(resolved_p)
                if has_txt:
                    cleaned_sections = [clean_page_text(s) for s in sections]
                    filtered_sections = filter_repetitive_headers_footers(cleaned_sections)
                    attachment_clean_text = postprocess_legal_text("\n\n".join(filtered_sections))
                extraction_error = err

        full_text = combine_feedback_and_attachment(raw_web_feedback, attachment_clean_text)
        has_text = len(full_text.strip()) > 10

        if has_text:
            stats["with_text"] += 1

        legal_refs = extract_legal_references(full_text)
        stats["articles_found"] += len(legal_refs["detected_articles"])
        stats["annexes_found"] += len(legal_refs["detected_annexes"])

        chunks = chunk_legal_text(full_text, target_words=target_chunk_words, overlap_words=overlap_words)
        stats["chunks_generated"] += len(chunks)

        parsed_record = {
            "feedback_id": feedback_id,
            "stage_name": conf["stage_name"],
            "publication_id": conf["publication_id"],
            "organization": org_name,
            "user_type": user_type,
            "country": country,
            "raw_web_feedback": raw_web_feedback,
            "attachment_name": attachment_name,
            "has_attachment": has_attachment_decl,
            "has_attachment_file": has_attachment_file,
            "has_text": has_text,
            "full_text": full_text,
            "total_chars": len(full_text),
            "total_words": len(full_text.split()),
            "detected_articles": legal_refs["detected_articles"],
            "detected_annexes": legal_refs["detected_annexes"],
            "detected_recitals": legal_refs["detected_recitals"],
            "chunks": chunks
        }
        if extraction_error:
            parsed_record["extraction_warning"] = extraction_error

        parsed_records.append(parsed_record)

        if not TQDM_AVAILABLE and (idx % 25 == 0 or idx == len(records)):
            print(f"[{idx}/{len(records)}] '{org_name}' | Parole: {len(full_text.split())} | Chunks: {len(chunks)}")

    os.makedirs(os.path.dirname(parsed_path), exist_ok=True)
    with open(parsed_path, "w", encoding="utf-8") as f:
        json.dump(parsed_records, f, ensure_ascii=False, indent=2)

    print("\n" + "-" * 50)
    print(f"🎉 COMPLETATO PARSING {conf['title']}!")
    print(f"💾 File salvato: '{conf['parsed_json']}'")
    print(f"📊 Statistiche:")
    print(f"   • Record totali:              {stats['total']}")
    print(f"   • Allegati estratti:          {stats['attachments_extracted']}")
    print(f"   • Record con testo valido:    {stats['with_text']}")
    print(f"   • Chunks generati:            {stats['chunks_generated']}")
    print(f"   • Riferimenti ad Articoli:    {stats['articles_found']}")
    print(f"   • Riferimenti ad Allegati:    {stats['annexes_found']}")
    print("-" * 50 + "\n")


# =====================================================================
# CLI
# =====================================================================

def main():
    parser = argparse.ArgumentParser(
        description="Estrae, pulisce e suddivide in chunk i feedback Have Your Say per Stage 1 e Stage 2."
    )
    parser.add_argument(
        "--procedure", "-p",
        default="ai_act",
        help="ID procedura in data_sources/procedures/<id>.yaml (default: ai_act)"
    )
    parser.add_argument(
        "--stage", "-s",
        choices=["1", "2", "all"],
        default="all",
        help="Quale stage elaborare: '1', '2' o 'all' (default: all)"
    )
    parser.add_argument("--chunk-size", "-c", type=int, default=400, help="Target parole per chunk (default: 400)")
    parser.add_argument("--chunk-overlap", type=int, default=50, help="Overlap parole (default: 50)")

    args = parser.parse_args()
    script_dir = os.path.dirname(os.path.abspath(__file__))
    stages_config = load_stages_config(args.procedure, script_dir)

    if args.stage in ["1", "all"]:
        if 1 in stages_config:
            parse_stage(1, script_dir, args.chunk_size, args.chunk_overlap, stages_config)
        elif args.stage == "1":
            print(f"⚠️ Stage 1 non configurato per la procedura '{args.procedure}' (publication è null).")

    if args.stage in ["2", "all"]:
        if 2 in stages_config:
            parse_stage(2, script_dir, args.chunk_size, args.chunk_overlap, stages_config)
        elif args.stage == "2":
            print(f"⚠️ Stage 2 non configurato per la procedura '{args.procedure}' (publication è null).")


if __name__ == "__main__":
    main()
